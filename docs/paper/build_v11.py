"""Build CHOLERA_PAPER_V11_GEE.docx from V10: refreshed screenshots from the live deployment,
new Figure 8b (verified state-tier choropleth + drill-down), prose/captions updated for the
national-tier UI and the removed LGA case panel / time-lapse. Figure 13 (benchmark) unchanged."""
import copy, re
import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

V10="/root/cholera_hod_tasks/CHOLERA_PAPER_V10_GEE.docx"
OUT="/root/cholera_hod_tasks/CHOLERA_PAPER_V11_GEE.docx"
F="/root/cholera_hod_tasks/figs_v11"

d=docx.Document(V10)
P=d.paragraphs
def find(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix): return p
    raise KeyError(prefix)
def set_text(p, text):
    r=p.runs
    if not r: p.add_run(text); return
    r[0].text=text
    for x in r[1:]: x._r.getparent().remove(x._r)
def img_par(caption_prefix):
    """Paragraph holding the picture immediately above the given caption."""
    for i,p in enumerate(d.paragraphs):
        if p.text.startswith(caption_prefix):
            for j in range(i-1, max(i-4,0), -1):
                if re.search(r'r:embed="rId\d+"', d.paragraphs[j]._p.xml): return d.paragraphs[j]
    raise KeyError(caption_prefix)
def swap_image(caption_prefix, png):
    p=img_par(caption_prefix)
    rid=re.search(r'r:embed="(rId\d+)"', p._p.xml).group(1)
    d.part.rels[rid].target_part._blob=open(png,"rb").read()
    return rid

# ── 1. swap screenshots ─────────────────────────────────────────────
swap_image("Figure 4:",  f"{F}/fig4_dashboard_national_tier.png")
swap_image("Figure 5:",  f"{F}/fig5_map_view.png")
swap_image("Figure 9:",  f"{F}/fig9_heuristic_tier_riskmap.png")
swap_image("Figure 10:", f"{F}/fig10_alerts.png")
swap_image("Figure 11:", f"{F}/fig11_copilot_live.png")
swap_image("Figure 12:", f"{F}/fig12_agent_explorer.png")

# ── 2. captions ─────────────────────────────────────────────────────
set_text(find("Figure 4:"),
 "Figure 4: National dashboard as deployed. The upper section is the verified national tier — a state-level choropleth of the NCDC situation-report extraction with the official NCDC year-end figures for 2021 (111,062 suspected cases; 3,604 deaths) shown alongside, and clearly separated from, the dataset-derived lower bound (98,531 cases across 33 reporting states). The lower section is the 774-LGA environmental risk map, labelled on screen as a heuristic tier.")
set_text(find("Figure 5:"),
 "Figure 5: Full-screen national map view for LGA-level exploration of the environmental risk layer and health-facility registry across the Federal Republic of Nigeria. Following exclusion of the national LGA-month case panel (Section 3.1.1), the LGA quick-stats no longer display case counts; an explicit evidence-tier note is shown in their place.")
set_text(find("Figure 9:"),
 "Figure 9: The heuristic LGA tier and its exploratory flood–cholera lag panel as rendered beneath the verified national tier. The on-screen label states that the risk score is derived from flood-archive and satellite inputs only and is not validated against outbreak data. Dependence and multiplicity are not corrected; no causal or predictive claim is made.")
set_text(find("Figure 10:"),
 "Figure 10: Alerts & Notifications page. With the LGA-month case panel excluded, the rule engine's case-surge triggers have no input and the page correctly reports no active alerts; flood-exposure and rainfall-threshold rules remain armed. The engine's epidemiological performance requires prospective evaluation and is not claimed here.")
set_text(find("Figure 11:"),
 "Figure 11: AI Surveillance Copilot (sidebar) answering a live query on the national dashboard, served through NVIDIA NIM (nemotron-3-super-120b-a12b), the same model used in the Section 6.2 benchmark. The assistant reports the top verified 2021 state burdens from the year-end view, quotes the official NCDC total separately, and labels the evidence tier of each figure. The provider status panel shows the single live provider; all others are in mock mode.")
set_text(find("Figure 12:"),
 "Figure 12: Agent Explorer dedicated page — copilot-assisted interactive dashboard workspace, demonstrated with the Cross River 2021 sentinel pilot line-list aggregate.")

# ── 3. §4.3 prose: dashboard now shows the national tier first; case panel gone; time-lapse removed
p=find("The dashboard and map successfully rendered geospatial risk layers")
set_text(p,
 "The dashboard and map render the two evidence tiers in the order the paper reports them. The upper section of the national dashboard (Figure 4) is the verified state tier of Section 5: a state choropleth of the NCDC situation-report extraction with a year selector, official-versus-dataset KPI cards, a year-by-year burden chart, and a per-state drill-down that lists every verified record with its epi-week, confidence tier, extraction method, monotonicity flag and a link to the source PDF (Figure 8b). Beneath it sits the 774-LGA environmental risk map (Figures 5 and 9), labelled on screen as a heuristic tier. Figure 6 shows the pilot case and death distribution derived directly from the ingested line-list. Consistent with the exclusion of the national LGA-month panel described in Section 3.1.1, the deployed interface displays no LGA-level case counts outside the four sentinel LGAs: the corresponding cards show an explicit evidence-tier note rather than a zero, and a 90-day time-lapse control that had animated the excluded panel was removed. The map component is built against the full 774-LGA GRID3 boundary set; unpopulated units are rendered as no-data rather than as zero risk.")

# ── 4. §5: new Figure 8b (state choropleth + drill-down) after Figure 8 caption
cap8=find("Figure 8:")
# picture paragraph clone (copy Figure 8's picture paragraph, then replace its image part with a NEW part)
pic8=img_par("Figure 8:")
new_pic=copy.deepcopy(pic8._p); cap8._p.addnext(new_pic)
new_pic_par=Paragraph(new_pic, cap8._parent)
# add a fresh image relationship so we don't clobber Figure 8's blob
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.parts.image import ImagePart
from docx.image.image import Image as DocxImage
img=DocxImage.from_file(f"{F}/fig4b_state_drilldown.png")
partname=d.part.package.next_partname("/word/media/image%d.png")
img_part=ImagePart(partname, img.content_type, img.blob, img)
new_rid=d.part.relate_to(img_part, RT.IMAGE)
xml=new_pic_par._p.xml
old_rid=re.search(r'r:embed="(rId\d+)"', xml).group(1)
for blip in new_pic_par._p.iter():
    if blip.tag.endswith('}blip'):
        for k in list(blip.attrib):
            if k.endswith('}embed'): blip.set(k, new_rid)
# caption for 8b
cap8b=copy.deepcopy(cap8._p); new_pic.addnext(cap8b)
cap8b_par=Paragraph(cap8b, cap8._parent)
set_text(cap8b_par,
 "Figure 8b: Verified national tier as deployed — state choropleth of 2021 year-end snapshots (33 reporting states; four states with no parsed report shown as no-data) and the per-state drill-down, here for Niger State, listing every verified record with epi-week, cumulative cases and deaths, CFR, confidence tier, extraction method, monotonicity flag and a link to the exact NCDC situation-report PDF. The dataset-derived national sum (98,531) is displayed as a lower bound beneath the official year-end figure (111,062), never merged with it.")
# reference sentence in §5 body
p=find("Officially reported cholera figures were compiled from NCDC situation reports (Table 7).")
set_text(p, p.text + " The same verified state-level series drives the national tier of the deployed dashboard (Figure 8b), which exposes every record's source PDF so that any figure on screen can be traced to the situation report it was extracted from.")

# ── 5. §6.1 alerts paragraph: reflect empty case input honestly
p=find("Beyond the epidemiological demonstrations, the platform provides a configurable rule-based alert engine")
set_text(p,
 "Beyond the epidemiological demonstrations, the platform provides a configurable rule-based alert engine (Figure 10) supporting case-surge, high-risk-score and recent-flooding triggers, together with PDF and CSV report export. In the deployment reported here the case-surge trigger has no input, because the national LGA-month case panel was excluded (Section 3.1.1) and verified case data exist only at state resolution; the flood-exposure and rainfall rules remain active. The alert engine's epidemiological performance — alert sensitivity, specificity and timeliness against observed outbreaks — requires prospective evaluation and is not claimed here.")

# ── 6. §6 copilot paragraph: note the year-end view + semantics guard
p=find("Figure 11 shows the AI Surveillance Copilot sidebar open")
set_text(p,
 "Figure 11 shows the AI Surveillance Copilot sidebar open on the national dashboard during a live query. Figure 12 shows the dedicated Agent Explorer page used for copilot-assisted interactive dashboards. Because the verified state series is cumulative by epi-week, a naive SQL aggregation over it inflates yearly totals several-fold; the deployed agent is therefore given an explicit data-semantics preamble and a database view (state_cholera_year_end) that exposes one year-end snapshot per state-year, and is instructed to label the evidence tier of every number it reports. The Cross River 2021 pilot aggregate is used in Figure 12 as a grounded demonstration dataset; the national architecture remains 774 LGAs.")

d.save(OUT)
print("saved", OUT)
